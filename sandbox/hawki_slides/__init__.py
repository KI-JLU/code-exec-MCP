"""
hawki_slides - a deck in a dozen calls, on the JLU template or a style of its own.

Why this exists: models write pptxgenjs and python-pptx from memory and get
them wrong on the first try - methods that do not exist, coordinates that
overlap, decks that never leave /tmp. OpenAI's sandbox mounts a helper module
for the same reason; this is HAWKI's.

    from hawki_slides import Deck

    deck = Deck(title="Anthropomorphisierung von LLM", author="HAWKI", lang="de")
    deck.title("Anthropomorphisierung von LLM", "Warum wir Sprachmodelle vermenschlichen")
    deck.bullets("Was bedeutet das?", ["Punkt eins", {"text": "Punkt zwei", "sub": ["Detail"]}],
                 sources=["https://..."])
    deck.cards("Vier Signale", [{"heading": "Dialog", "text": "..."}, {"heading": "Ich-Perspektive", "text": "..."}])
    deck.two_columns("Hilfreich vs. irreführend", {"heading": "Hilfreich", "items": ["..."]},
                                                  {"heading": "Irreführend", "items": ["..."]})
    deck.quote("Menschlich genug, um zu helfen.", "Fazit")
    deck.closing("Danke · Fragen?")
    deck.save("/tmp/anthropomorphisierung.pptx")

Where the look comes from, in this order:
  - template="attached"                   the .potx or .pptx the user attached
    (or template="/work/<name>" to name one of several)
  - style="jlu" (the default)             the JLU corporate template, German or
                                          English by `lang`
  - style="purple" | "blue" | "green" | "red" | "slate"
                                          HAWKI's own drawn styles, no template

Every slide method accepts notes="..." and sources=[...]; both land in the
speaker notes, the sources as a "[Sources]" list. Every method returns the
Deck, so calls chain.
"""

from __future__ import annotations

import datetime as _dt
import io
import os
import re
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

__all__ = ["Deck", "STYLES", "TEMPLATES", "template_for"]

TEMPLATE_DIR = os.environ.get("HAWKI_SLIDES_TEMPLATES", "/usr/local/share/hawki-slides/templates")

TEMPLATES = {
    "de": os.path.join(TEMPLATE_DIR, "JLU-de.potx"),
    "en": os.path.join(TEMPLATE_DIR, "JLU-en.potx"),
}

STYLES = {
    "purple": dict(accent="6C5CE7", ink="1D2433", muted="667085", pale="EEF0FF", line="D8DCE8", paper="F7F8FC", dark="101525"),
    "blue":   dict(accent="2563EB", ink="0F172A", muted="64748B", pale="E8EEFF", line="D6DCE8", paper="F6F8FC", dark="0B1220"),
    "green":  dict(accent="0F8A5F", ink="13221B", muted="5E6E66", pale="E6F5EE", line="D2E2DA", paper="F5FAF7", dark="0B1A13"),
    "red":    dict(accent="D9483B", ink="23181A", muted="6F6366", pale="FBECEA", line="E7D6D4", paper="FBF7F6", dark="1C0F10"),
    "slate":  dict(accent="334155", ink="0F172A", muted="64748B", pale="EEF2F7", line="D9DFE7", paper="F7F9FB", dark="0B1220"),
}

# Slide geometry of LAYOUT_WIDE, the size both the JLU template and the drawn
# styles use. Content boxes derive from these, never from literals in methods.
W, H = 13.333, 7.5
MARGIN = 0.93
CONTENT_TOP = 1.95
CONTENT_BOTTOM = 6.85


ATTACHED_DIR = "/work"


def attached_templates() -> list:
    """The .potx/.pptx files the user attached to the request, as HAWKI placed them in /work."""
    try:
        names = sorted(os.listdir(ATTACHED_DIR))
    except OSError:
        return []
    return [os.path.join(ATTACHED_DIR, n) for n in names if n.lower().endswith((".potx", ".pptx")) and n != "code.py"]


def _resolve_template(template) -> str:
    """
    template="attached" means: whatever .potx/.pptx the user attached. A bare
    file name is looked up in /work. A path is taken as it is.
    """
    value = _s(template)
    if value.lower() in ("attached", "upload", "uploaded", "user"):
        found = attached_templates()
        if not found:
            raise FileNotFoundError(
                "hawki_slides: no attached template - nothing ending in .potx or .pptx in %s (contents: %s). "
                "Leave `template` out to use the JLU template." % (ATTACHED_DIR, ", ".join(sorted(os.listdir(ATTACHED_DIR))) or "empty")
            )
        return found[0]
    if os.path.exists(value):
        return value
    candidate = os.path.join(ATTACHED_DIR, os.path.basename(value))
    if os.path.exists(candidate):
        return candidate
    raise FileNotFoundError(
        "hawki_slides: template not found: %r. Attached files are in %s: %s" % (
            template, ATTACHED_DIR, ", ".join(sorted(os.listdir(ATTACHED_DIR))) if os.path.isdir(ATTACHED_DIR) else "-"
        )
    )


def template_for(lang: str | None) -> str:
    """The JLU template for a language: English for en*, German for everything else."""
    return TEMPLATES["en"] if str(lang or "de").lower().startswith("en") else TEMPLATES["de"]


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.replace("#", "").upper())


def _s(value) -> str:
    return "" if value is None else str(value)


def _items(value) -> list:
    if value is None:
        return []
    return list(value) if isinstance(value, (list, tuple)) else [value]


def _fit_font_size(lines, box_h_in, box_w_in, largest=20, smallest=12):
    """
    The font size a list of lines fits in: steps down as the text grows so a
    slide with ten long bullets does not run off the bottom. Measured for a
    box of the given size with ~1.3 line spacing; PowerPoint's own autofit is
    set too, but only applies after the file is edited, and LibreOffice - which
    renders the previews - does not apply it at all.
    """
    size = largest
    while size > smallest:
        chars_per_line = max(10, int((box_w_in * 72) / (size * 0.5)))
        wrapped = sum(max(1, -(-len(_s(line)) // chars_per_line)) for line in lines)
        needed = (wrapped * size * 1.3 + len(lines) * size * 0.4) / 72
        if needed <= box_h_in:
            break
        size -= 1
    return size


def _open_template(path: str) -> Presentation:
    """
    Opens a .potx or .pptx as a presentation and removes every slide it came
    with. A .potx declares the template content type, which python-pptx does
    not open; the package is the same otherwise, so the type is rewritten in
    memory. Layouts, masters and theme survive - that is what a template is for.
    """
    with open(path, "rb") as handle:
        raw = handle.read()

    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                data = data.replace(b"presentationml.slideshow.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    buf.seek(0)

    prs = Presentation(buf)

    # The brand marks first. A corporate template often carries its logo as a
    # picture ON the sample slides rather than on the layout - the JLU one
    # does: the logo block top right of the title slide is 'Grafik 3' on
    # sample slide 1, and nowhere else. So before the samples go, every
    # non-placeholder picture on the first sample of each layout is kept, with
    # its position, and put back on the slides this deck adds on that layout.
    brand_pictures = {}
    for sample in prs.slides:
        layout_name = sample.slide_layout.name
        if layout_name in brand_pictures:
            continue
        pictures = []
        for shape in sample.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and not shape.is_placeholder:
                try:
                    pictures.append((shape.image.blob, shape.left, shape.top, shape.width, shape.height))
                except Exception:
                    pass
        brand_pictures[layout_name] = pictures

    # Drop the sample slides: the deck starts empty, on the template's layouts.
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)

    prs._hawki_brand_pictures = brand_pictures
    return prs


def _placeholders(layout, kind):
    return [ph for ph in layout.placeholders if ph.placeholder_format.type == kind]


def _find_layout(prs, names, predicate):
    """A layout by any of its expected names, else the first one the predicate accepts."""
    by_name = {layout.name.strip().lower(): layout for layout in prs.slide_layouts}
    for name in names:
        if name.lower() in by_name:
            return by_name[name.lower()]
    for layout in prs.slide_layouts:
        if predicate(layout):
            return layout
    return None


class Deck:
    def __init__(self, title="", author="HAWKI", lang="de", style="jlu", template=None, date=None):
        self.deck_title = _s(title)
        self.author = _s(author)
        self.lang = _s(lang) or "de"
        self.date = date or _dt.date.today().strftime("%d.%m.%Y" if not self.lang.lower().startswith("en") else "%Y-%m-%d")
        self.slide_count = 0

        if template:
            template = _resolve_template(template)
            self.mode = "template"
            self.template_path = template
        elif str(style).lower() == "jlu":
            self.mode = "template"
            self.template_path = template_for(self.lang)
        elif str(style).lower() in STYLES:
            self.mode = "drawn"
            self.colors = STYLES[str(style).lower()]
        else:
            raise ValueError("hawki_slides: unknown style %r - one of jlu, %s" % (style, ", ".join(STYLES)))

        if self.mode == "template":
            self.prs = _open_template(self.template_path)
            self._resolve_layouts()
        else:
            self.prs = Presentation()
            self.prs.slide_width = Inches(W)
            self.prs.slide_height = Inches(H)

    # ------------------------------------------------------------------ layouts

    def _resolve_layouts(self):
        prs = self.prs
        has = lambda layout, kind, n=1: len(_placeholders(layout, kind)) >= n  # noqa: E731

        self.L = {
            "title": _find_layout(prs, ["Titel m. Headline und Subline", "Title Slide", "Titelfolie"],
                                  lambda l: has(l, PP_PLACEHOLDER.CENTER_TITLE) and has(l, PP_PLACEHOLDER.BODY))
                     or _find_layout(prs, [], lambda l: has(l, PP_PLACEHOLDER.CENTER_TITLE))
                     or prs.slide_layouts[0],
            "text": _find_layout(prs, ["1_Textfolie nur Titel und Text", "Textfolie nur Titel und Text", "Title and Content", "Titel und Inhalt"],
                                 lambda l: has(l, PP_PLACEHOLDER.TITLE) and len(_placeholders(l, PP_PLACEHOLDER.BODY)) == 1)
                    or _find_layout(prs, [], lambda l: has(l, PP_PLACEHOLDER.TITLE) and has(l, PP_PLACEHOLDER.BODY))
                    or prs.slide_layouts[min(1, len(prs.slide_layouts) - 1)],
            "quote": _find_layout(prs, ["Zitat", "Quote"],
                                  lambda l: has(l, PP_PLACEHOLDER.TITLE) and len(_placeholders(l, PP_PLACEHOLDER.BODY)) == 2
                                  and "zitat" in l.name.lower()),
            "closing": _find_layout(prs, ["Ende", "End", "Schluss", "Closing"],
                                    lambda l: has(l, PP_PLACEHOLDER.CENTER_TITLE) and not has(l, PP_PLACEHOLDER.BODY)
                                    and not has(l, PP_PLACEHOLDER.PICTURE)),
            "image": _find_layout(prs, ["Textfolie nur Titel und Bild", "Title and Picture", "Titel und Bild"],
                                  lambda l: has(l, PP_PLACEHOLDER.TITLE) and has(l, PP_PLACEHOLDER.PICTURE)),
        }

    def _new_slide(self, role):
        self.slide_count += 1
        layout = self.L.get(role) or self.L["text"]
        slide = self.prs.slides.add_slide(layout)
        for blob, left, top, width, height in getattr(self.prs, "_hawki_brand_pictures", {}).get(layout.name, []):
            slide.shapes.add_picture(io.BytesIO(blob), left, top, width, height)
        return slide

    # ------------------------------------------------------------------ helpers

    @staticmethod
    def _ph(slide, kind, index=0):
        found = [ph for ph in slide.placeholders if ph.placeholder_format.type == kind]
        return found[index] if len(found) > index else None

    @staticmethod
    def _remove(shape):
        if shape is not None:
            shape._element.getparent().remove(shape._element)

    def _remove_unfilled_placeholders(self, slide):
        for ph in list(slide.placeholders):
            kind = ph.placeholder_format.type
            if kind in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                continue
            if kind == PP_PLACEHOLDER.PICTURE:
                self._remove(ph)
            elif ph.has_text_frame and not ph.text_frame.text.strip():
                self._remove(ph)

    def _notes(self, slide, notes=None, sources=None):
        parts = []
        if notes:
            parts.append(_s(notes))
        srcs = [_s(x) for x in _items(sources) if _s(x)]
        if srcs:
            parts.append("[Sources]\n" + "\n".join("- " + x for x in srcs))
        if parts:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    @staticmethod
    def _set_text(shape, text, size=None, bold=None, color=None, align=None):
        tf = shape.text_frame
        tf.text = _s(text)
        for paragraph in tf.paragraphs:
            if align is not None:
                paragraph.alignment = align
            for run in paragraph.runs:
                if size is not None:
                    run.font.size = Pt(size)
                if bold is not None:
                    run.font.bold = bold
                if color is not None:
                    run.font.color.rgb = _rgb(color)

    @staticmethod
    def _fill_bullets(text_frame, items, size, color=None):
        """Bullets into a text frame: strings, or {"text", "sub": [...]} for a second level."""
        text_frame.clear()
        first = True
        for item in _items(items):
            text = _s(item.get("text")) if isinstance(item, dict) else _s(item)
            subs = _items(item.get("sub")) if isinstance(item, dict) else []
            for level, line in [(0, text)] + [(1, _s(sub)) for sub in subs]:
                paragraph = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
                first = False
                paragraph.text = line
                paragraph.level = level
                paragraph.space_after = Pt(6 if level == 0 else 3)
                for run in paragraph.runs:
                    run.font.size = Pt(size if level == 0 else max(size - 3, 10))
                    if color is not None:
                        run.font.color.rgb = _rgb(color)
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    @staticmethod
    def _lines_of(items):
        lines = []
        for item in _items(items):
            if isinstance(item, dict):
                lines.append(_s(item.get("text")))
                lines.extend("    " + _s(s) for s in _items(item.get("sub")))
            else:
                lines.append(_s(item))
        return lines

    def _textbox(self, slide, x, y, w, h, text="", size=14, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        self._set_text(box, text, size=size, bold=bold, color=color, align=align)
        return box

    def _rect(self, slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        rect = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = _rgb(fill)
        if line:
            rect.line.color.rgb = _rgb(line)
            rect.line.width = Pt(0.75)
        else:
            rect.line.fill.background()
        rect.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            rect.adjustments[0] = 0.06
        rect.text_frame.text = ""
        return rect

    def _content_area(self, slide):
        """Where the body goes: the template's body placeholder if there is one, else the default box."""
        body = self._ph(slide, PP_PLACEHOLDER.BODY)
        if body is not None:
            return (body.left / 914400, body.top / 914400, body.width / 914400, body.height / 914400)
        return (MARGIN, CONTENT_TOP, W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP)

    def _accent(self):
        if self.mode == "drawn":
            return self.colors["accent"]
        return "156082"  # the JLU blue, and a sane default for an unknown template

    # ------------------------------------------------------------- drawn chrome

    def _drawn_frame(self, slide, dark=False):
        c = self.colors
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(c["dark"] if dark else c["paper"])
        if not dark:
            self._rect(slide, 0, 0, W, 0.16, c["accent"], shape=MSO_SHAPE.RECTANGLE)
            self._textbox(slide, MARGIN, 7.06, 7, 0.25, self.deck_title.upper(), size=8, color=c["muted"])
            right = " · ".join(x for x in [self.author, str(_dt.date.today().year)] if x)
            self._textbox(slide, W - MARGIN - 3.4, 7.06, 2.9, 0.25, right, size=8, color=c["muted"], align=PP_ALIGN.RIGHT)
            self._textbox(slide, W - MARGIN - 0.5, 7.06, 0.5, 0.25, str(self.slide_count), size=8, color=c["muted"], align=PP_ALIGN.RIGHT)

    def _drawn_heading(self, slide, title, subtitle=None):
        c = self.colors
        self._textbox(slide, MARGIN, 0.5, W - 2 * MARGIN, 0.7, title, size=28, bold=True, color=c["ink"])
        if subtitle:
            self._textbox(slide, MARGIN, 1.15, W - 2 * MARGIN, 0.4, subtitle, size=13, color=c["muted"])

    def _drawn_slide(self, dark=False):
        self.slide_count += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank
        self._drawn_frame(slide, dark=dark)
        return slide

    # ----------------------------------------------------------------- slides

    def title(self, title, subtitle=None, notes=None, sources=None):
        if self.mode == "drawn":
            c = self.colors
            slide = self._drawn_slide(dark=True)
            self._rect(slide, 0, 0, 0.3, H, c["accent"], shape=MSO_SHAPE.RECTANGLE)
            self._textbox(slide, 1.1, 2.1, W - 2.2, 1.6, title, size=40, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.BOTTOM)
            if subtitle:
                self._textbox(slide, 1.1, 3.85, W - 2.2, 1.0, subtitle, size=18, color="D7DDF0")
            meta = " · ".join(x for x in [self.author, self.date] if x)
            self._textbox(slide, 1.1, 6.6, 8, 0.3, meta, size=10, color="8E97B6")
        else:
            slide = self._new_slide("title")
            head = self._ph(slide, PP_PLACEHOLDER.CENTER_TITLE) or self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title)
            sub = self._ph(slide, PP_PLACEHOLDER.BODY)
            if sub is not None:
                sub.text_frame.text = _s(subtitle) if subtitle else " · ".join(x for x in [self.author, self.date] if x)
            self._remove_unfilled_placeholders(slide)
        self._notes(slide, notes, sources)
        return self

    def bullets(self, title, items, subtitle=None, notes=None, sources=None):
        lines = self._lines_of(items)
        if self.mode == "drawn":
            c = self.colors
            slide = self._drawn_slide()
            self._drawn_heading(slide, title, subtitle)
            x, y, w, h = MARGIN, CONTENT_TOP, W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            self._fill_bullets(box.text_frame, items, _fit_font_size(lines, h, w, 20, 12), color=c["ink"])
            for paragraph in box.text_frame.paragraphs:
                # a textbox has no bullet style of its own; borrow one
                pPr = paragraph._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(0.3 + 0.35 * paragraph.level))))
                pPr.set("indent", str(int(-Inches(0.25))))
                for child in list(pPr):
                    if child.tag.endswith("}buNone"):
                        pPr.remove(child)
                bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
                pPr.append(bu)
        else:
            slide = self._new_slide("text")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title)
            body = self._ph(slide, PP_PLACEHOLDER.BODY)
            if body is None:
                x, y, w, h = self._content_area(slide)
                body = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            bw, bh = body.width / 914400, body.height / 914400
            size = _fit_font_size(lines, bh - 0.2, bw - 0.3, 20, 12)
            self._fill_bullets(body.text_frame, items, size)
            if subtitle:
                # the JLU text layouts have no subtitle box; the first paragraph takes it
                first = body.text_frame.paragraphs[0]
                new = first._p.makeelement(first._p.tag, {})
                first._p.addprevious(new)
                from pptx.text.text import _Paragraph
                para = _Paragraph(new, body.text_frame)
                para.text = _s(subtitle)
                for run in para.runs:
                    run.font.size = Pt(size)
                    run.font.bold = True
                pPr = para._p.get_or_add_pPr()
                pPr.append(pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", {}))
            self._remove_unfilled_placeholders(slide)
        self._notes(slide, notes, sources)
        return self

    def cards(self, title, cards, subtitle=None, notes=None, sources=None):
        cards = _items(cards)[:6]
        n = max(1, len(cards))
        cols = n if n <= 3 else -(-n // 2)
        rows = -(-n // cols)
        accent = self._accent()

        if self.mode == "drawn":
            slide = self._drawn_slide()
            self._drawn_heading(slide, title, subtitle)
            x0, y0, aw, ah = MARGIN, CONTENT_TOP, W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
            fill, line, ink, muted = "FFFFFF", self.colors["line"], self.colors["ink"], self.colors["muted"]
        else:
            slide = self._new_slide("text")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title)
            x0, y0, aw, ah = self._content_area(slide)
            self._remove(self._ph(slide, PP_PLACEHOLDER.BODY))
            self._remove_unfilled_placeholders(slide)
            fill, line, ink, muted = "FFFFFF", "D9DEE6", "1D2433", "5B6472"

        gap = 0.3
        cw = (aw - gap * (cols - 1)) / cols
        ch = (ah - gap * (rows - 1)) / rows
        for i, card in enumerate(cards):
            cx = x0 + (i % cols) * (cw + gap)
            cy = y0 + (i // cols) * (ch + gap)
            self._rect(slide, cx, cy, cw, ch, fill, line=line)
            circle = self._rect(slide, cx + 0.25, cy + 0.25, 0.45, 0.45, accent, shape=MSO_SHAPE.OVAL)
            self._set_text(circle, str(i + 1), size=12, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
            circle.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            heading = _s((card or {}).get("heading") if isinstance(card, dict) else card)
            body = _s((card or {}).get("text")) if isinstance(card, dict) else ""
            self._textbox(slide, cx + 0.85, cy + 0.22, cw - 1.1, 0.5, heading, size=16, bold=True, color=ink)
            self._textbox(slide, cx + 0.25, cy + 0.85, cw - 0.5, ch - 1.05, body,
                          size=_fit_font_size([body], ch - 1.05, cw - 0.5, 14, 10), color=muted)
        self._notes(slide, notes, sources)
        return self

    def two_columns(self, title, left, right, subtitle=None, notes=None, sources=None):
        accent = self._accent()
        if self.mode == "drawn":
            slide = self._drawn_slide()
            self._drawn_heading(slide, title, subtitle)
            x0, y0, aw, ah = MARGIN, CONTENT_TOP, W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
            fills, line, ink = (self.colors["pale"], "FFFFFF"), self.colors["line"], self.colors["ink"]
        else:
            slide = self._new_slide("text")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title)
            x0, y0, aw, ah = self._content_area(slide)
            self._remove(self._ph(slide, PP_PLACEHOLDER.BODY))
            self._remove_unfilled_placeholders(slide)
            fills, line, ink = ("F1F5F9", "FFFFFF"), "D9DEE6", "1D2433"

        gap = 0.4
        cw = (aw - gap) / 2
        for col, cx, fill in ((left, x0, fills[0]), (right, x0 + cw + gap, fills[1])):
            col = col if isinstance(col, dict) else {"heading": "", "items": col}
            items = _items(col.get("items"))
            self._rect(slide, cx, y0, cw, ah, fill, line=line)
            self._textbox(slide, cx + 0.3, y0 + 0.2, cw - 0.6, 0.5, _s(col.get("heading")), size=18, bold=True, color=accent)
            if items:
                box = slide.shapes.add_textbox(Inches(cx + 0.3), Inches(y0 + 0.85), Inches(cw - 0.6), Inches(ah - 1.05))
                self._fill_bullets(box.text_frame, items, _fit_font_size(self._lines_of(items), ah - 1.05, cw - 0.6, 16, 11), color=ink)
                for paragraph in box.text_frame.paragraphs:
                    pPr = paragraph._p.get_or_add_pPr()
                    pPr.set("marL", str(int(Inches(0.28))))
                    pPr.set("indent", str(int(-Inches(0.22))))
                    pPr.append(pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"}))
        self._notes(slide, notes, sources)
        return self

    twoColumns = two_columns  # the JavaScript spelling, for a model that remembers that one

    def quote(self, text, attribution=None, title=None, notes=None, sources=None):
        if self.mode == "drawn":
            c = self.colors
            slide = self._drawn_slide()
            self._rect(slide, MARGIN, 1.6, W - 2 * MARGIN, 4.2, c["dark"])
            self._textbox(slide, MARGIN + 0.7, 2.0, W - 2 * MARGIN - 1.4, 2.8, text,
                          size=_fit_font_size([text], 2.8, W - 2 * MARGIN - 1.4, 34, 18), bold=True, color="FFFFFF",
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if attribution:
                self._textbox(slide, MARGIN + 0.7, 4.9, W - 2 * MARGIN - 1.4, 0.5, attribution, size=14, color="B8C0D5", align=PP_ALIGN.CENTER)
        elif self.L.get("quote") is not None:
            slide = self._new_slide("quote")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title or "")
            bodies = [ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.BODY]
            bodies.sort(key=lambda ph: ph.top)
            if bodies:
                bodies[0].text_frame.text = "»" + _s(text) + "«"
            if len(bodies) > 1 and attribution:
                bodies[1].text_frame.text = _s(attribution)
            self._remove_unfilled_placeholders(slide)
        else:
            slide = self._new_slide("text")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title or "")
            x0, y0, aw, ah = self._content_area(slide)
            self._remove(self._ph(slide, PP_PLACEHOLDER.BODY))
            self._remove_unfilled_placeholders(slide)
            self._textbox(slide, x0, y0 + 0.4, aw, ah - 1.4, "»" + _s(text) + "«", size=_fit_font_size([text], ah - 1.4, aw, 32, 18),
                          bold=True, color=self._accent(), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if attribution:
                self._textbox(slide, x0, y0 + ah - 0.9, aw, 0.6, attribution, size=14, color="5B6472", align=PP_ALIGN.CENTER)
        self._notes(slide, notes, sources)
        return self

    def closing(self, text="Danke", subtitle=None, notes=None, sources=None):
        if self.mode == "drawn":
            c = self.colors
            slide = self._drawn_slide()
            self._textbox(slide, MARGIN, 2.6, W - 2 * MARGIN, 1.2, text, size=36, bold=True, color=c["accent"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if subtitle:
                self._textbox(slide, MARGIN, 3.9, W - 2 * MARGIN, 0.8, subtitle, size=16, color=c["muted"], align=PP_ALIGN.CENTER)
        else:
            slide = self._new_slide("closing" if self.L.get("closing") is not None else "title")
            head = self._ph(slide, PP_PLACEHOLDER.CENTER_TITLE) or self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(text)
            sub = self._ph(slide, PP_PLACEHOLDER.BODY)
            if sub is not None:
                if subtitle:
                    sub.text_frame.text = _s(subtitle)
                else:
                    self._remove(sub)
            self._remove_unfilled_placeholders(slide)
        self._notes(slide, notes, sources)
        return self

    def image(self, title, path, caption=None, notes=None, sources=None):
        if not os.path.exists(_s(path)):
            raise FileNotFoundError("hawki_slides: image not found: %r" % path)
        if self.mode == "drawn":
            slide = self._drawn_slide()
            self._drawn_heading(slide, title)
            x0, y0, aw, ah = MARGIN, CONTENT_TOP, W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
            pic = slide.shapes.add_picture(path, Inches(x0), Inches(y0), height=Inches(ah - (0.5 if caption else 0)))
            if pic.width > Inches(aw):
                pic.width, pic.height = Inches(aw), int(pic.height * Inches(aw) / pic.width)
            pic.left = int(Inches(x0) + (Inches(aw) - pic.width) / 2)
            if caption:
                self._textbox(slide, x0, CONTENT_BOTTOM - 0.45, aw, 0.4, caption, size=11, color=self.colors["muted"], align=PP_ALIGN.CENTER)
        else:
            slide = self._new_slide("image" if self.L.get("image") is not None else "text")
            head = self._ph(slide, PP_PLACEHOLDER.TITLE)
            if head is not None:
                head.text_frame.text = _s(title)
            picture_ph = self._ph(slide, PP_PLACEHOLDER.PICTURE)
            if picture_ph is not None:
                picture_ph.insert_picture(path)
            else:
                x0, y0, aw, ah = self._content_area(slide)
                self._remove(self._ph(slide, PP_PLACEHOLDER.BODY))
                slide.shapes.add_picture(path, Inches(x0), Inches(y0), height=Inches(ah - 0.5))
            body = self._ph(slide, PP_PLACEHOLDER.BODY)
            if body is not None:
                if caption:
                    body.text_frame.text = _s(caption)
                else:
                    self._remove(body)
            self._remove_unfilled_placeholders(slide)
        self._notes(slide, notes, sources)
        return self

    # ------------------------------------------------------------------- save

    def save(self, path="/tmp/deck.pptx"):
        """Writes the deck. Under /tmp - the working directory is read-only. Returns the path."""
        path = _s(path) or "/tmp/deck.pptx"
        if not path.startswith("/tmp/"):
            raise ValueError("hawki_slides: save() needs a path under /tmp - the working directory is read-only (got %r)" % path)
        if not path.lower().endswith(".pptx"):
            path += ".pptx"
        self.prs.save(path)
        return path

    @property
    def raw(self):
        """The python-pptx Presentation, for anything this API does not cover."""
        return self.prs
